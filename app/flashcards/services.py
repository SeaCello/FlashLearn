from .models import UserFlashcard
from gpt.api import generate_flashcards
import fitz
import chardet
import docx
from fpdf import FPDF
from io import BytesIO

class FlashcardService:
    @staticmethod
    def extract_text_from_file(file):
        """
        Extract text from a file based on its extension.
        The supported extensions are: .pdf, .docx and .txt.
        """
        if file.name.lower().endswith(".pdf"):
            return FlashcardService._extract_text_from_pdf(file)
        elif file.name.lower().endswith(".docx"):
            return FlashcardService._extract_text_from_docx(file)
        else:
            return FlashcardService._extract_text_from_txt(file)
        
        
    @staticmethod
    def _extract_text_from_pdf(file):
        text = ""
        with fitz.open(stream=file.read(), filetype="pdf") as doc:
            for page in doc:
                text += page.get_text()
        return text[:6000]
    
    @staticmethod
    def _extract_text_from_docx(file):
        doc = docx.Document(file)
        return '\n'.join([p.text for p in doc.paragraphs])[:6000]
    
    @staticmethod
    def _extract_text_from_txt(file):
        raw_content = file.read()
        encoding = chardet.detect(raw_content).get('encoding')
        file.seek(0)
        return raw_content.decode(encoding)[:6000]
    
    @staticmethod
    def create_flashcards_from_file(user, file):
        text = FlashcardService.extract_text_from_file(file)
        api_flashcards = generate_flashcards(text)
        
        db_flashcards = []
        for data in api_flashcards:
            card = UserFlashcard.objects.create(
                user=user,
                title=data['title'],
                content=data['content']
            )
            db_flashcards.append({
                'id': card.id,
                'title': card.title,
                'content': card.content
            })
        
        return db_flashcards


    @staticmethod
    def update_flashcards_from_data(user, titles, contents, flashcard_ids):
        updated_flashcards = []
        
        for i in range(len(titles)):
            if i >= len(contents): 
                continue
                
            title = titles[i]
            content = contents[i]
            
            card_id = None
            if i < len(flashcard_ids) and flashcard_ids[i] and flashcard_ids[i].strip():
                try:
                    card_id = int(flashcard_ids[i])
                except (ValueError, TypeError):
                    card_id = None
            
            if card_id:
                card, _ = UserFlashcard.objects.update_or_create(
                    id=card_id, 
                    user=user,
                    defaults={'title': title, 
                                'content': content}
                )
            else:
                card = UserFlashcard.objects.create(
                    user=user, 
                    title=title, 
                    content=content
                )
            
            updated_flashcards.append({
                'id': card.id,
                'title': card.title, 
                'content': card.content
            })
        
        return updated_flashcards

class PDFService:
    @staticmethod
    def generate_flashcards_pdf(flashcards):
        if not flashcards:
            return None
        
        buffer_pdf = BytesIO()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, txt="Flashcards Gerados", ln=1, align='C')
        pdf.ln(10)
        
        
        pdf.set_font("Arial", size=12)
        for idx, card in enumerate(flashcards, 1):
            # Título
            pdf.set_font("Arial", 'B', 12)
            titulo = f"{idx}. {card.title}"
            safe_titulo = titulo.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_titulo)
            
            # Conteúdo
            pdf.set_font("Arial", '', 12)
            conteudo = f"{card.content}"
            safe_conteudo = conteudo.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_conteudo)
            
            pdf.ln(5)
        
        pdf_content = pdf.output(dest='S').encode('latin-1')
        buffer_pdf.write(pdf_content)
        buffer_pdf.seek(0)
        return buffer_pdf
    